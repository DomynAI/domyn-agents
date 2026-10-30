"""Check a .pptx for colliding text: words printed over each other, over a picture, cut by a
rule, hidden under a shape, spilling out of their box or card, or running off the slide.

    python check_pptx.py deck.pptx                 # a report; exit status 1 if anything collides
    python check_pptx.py deck.pptx --json          # the same, as JSON
    python check_pptx.py deck.pptx --render out/   # also out/<deck>.pdf and one PNG per slide

A slide made from a template's layout is checked with what the layout draws under it:
its background, cards, rules and fixed text, below the slide's own shapes.

python-pptx records where a text box is, not where its text ends up: text that does not fit
spills out of its box and PowerPoint draws it anyway. So this lays the text out itself, the
way PowerPoint does: every run measured with its own DM Sans file (`../../../resources/assets/
fonts/` from this file), wrapped at the box's width minus its insets and the paragraph's
indent, each line 1.2 x its size x the paragraph's line spacing, plus the space before and
after paragraphs, placed by the box's vertical anchor. A table row grows to fit its cells, as
it does in PowerPoint, and pushes the rows below it down. It then compares the ink of every
line (cap height to descender), not its line box, with everything else on the slide:

- text on text: lines of two boxes, or of one, printed over each other;
- text on a picture: a line over a picture that is not the slide's background (a chart, a
  screenshot);
- text crossed by a line: a rule drawn through the letters (a title wrapped onto the title
  rule, a card's body run into the card's divider);
- text out of its card: a line that is partly on a filled shape (a card, the table's header
  band) and partly off it;
- text under a shape: a line covered by a filled shape added after it;
- text overflowing its box: more lines than the box holds, or a word wider than the box;
- table row grown: a row taller than set because its text does not fit, which moves the
  rest of the table down;
- text off the slide, or outside the safe area (x 0.29 to 9.68 in on a 10 in slide).

`--render` also draws the deck with LibreOffice (`soffice`) into a PDF and, if pymupdf is
installed, one PNG per slide, to look at. LibreOffice is given the bundled DM Sans through a
private profile, so the render is right on a machine where DM Sans is not installed.

Needs python-pptx (which brings Pillow).
"""

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import asdict, dataclass, field
from functools import cache
from pathlib import Path
from typing import Any

from PIL import ImageFont
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Length

RESOURCES = Path(__file__).resolve().parents[3] / "resources"
FONTS = RESOURCES / "assets" / "fonts"
TYPOGRAPHY = json.loads((RESOURCES / "typography.json").read_text())
EMU = 914400  # per inch

# Single line spacing is 1.2 x the font size, in PowerPoint and in LibreOffice (measured on
# DM Sans); the font's own ascent + descent (1.302) is not what either uses.
LINE_PITCH = 1.2
# The first baseline sits this share of a line's pitch below the top of the line.
BASELINE = 1 / LINE_PITCH
# Where the ink is, relative to the baseline, as a share of the size: DM Sans's cap height
# above it, its descender below.
CAP_HEIGHT = 0.70
DESCENDER = 0.18
# Smaller overlaps, in inches, are rounding (0.01 in is 0.72 pt).
TOLERANCE = 0.01
# Text may run past the bottom of its box by this share of a line (a 23pt title needs 0.383
# in, its box is 0.38); more and a line has been added.
OVERFLOW_SHARE = 0.25
# The layouts' safe area, on a 10 in wide slide.
SAFE_X = (0.29, 9.68)
SAFE_TOLERANCE = 0.02
# A picture covering this share of the slide, as the first shape, is the background.
BACKGROUND_SHARE = 0.95
DEFAULT_SIZE = 18.0  # pt, PowerPoint's when a run sets none
DEFAULT_FAMILY = TYPOGRAPHY["styles"]["regular"]["family"]
# Families the template names but does not ship, measured with the nearest one shipped.
STAND_INS = {"DM Sans Thin": "DM Sans Light"}


@dataclass
class Line:
    """One line of text as laid out, measured by its ink, in inches."""

    slide: int
    shape: str
    order: int  # the shape's place in the z-order
    text: str
    x0: float
    x1: float
    baseline: float
    size: float  # pt, the largest on the line

    @property
    def top(self) -> float:
        return self.baseline - CAP_HEIGHT * self.size / 72

    @property
    def bottom(self) -> float:
        return self.baseline + DESCENDER * self.size / 72


@dataclass
class Rect:
    name: str
    order: int
    x0: float
    y0: float
    x1: float
    y1: float


@dataclass
class Segment:
    name: str
    x0: float
    y0: float
    x1: float
    y1: float


@dataclass
class Finding:
    slide: int
    kind: str
    detail: str
    where: tuple[float, float, float, float]  # inches: x0, y0, x1, y1


@dataclass
class Slide:
    """What the checks compare, for one slide."""

    number: int
    lines: list[Line] = field(default_factory=list)
    pictures: list[Rect] = field(default_factory=list)
    fills: list[Rect] = field(default_factory=list)
    segments: list[Segment] = field(default_factory=list)


@dataclass
class Report:
    findings: list[Finding] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    def note(self, text: str) -> None:
        if text not in self.notes:
            self.notes.append(text)


# --- fonts ------------------------------------------------------------------------------


@cache
def _font_files() -> dict[tuple[str, bool, bool], Path]:
    """(family, bold, italic) -> file, as typography.json pairs them. The family is the name
    PowerPoint stores in the deck, which is not always the name a font file reports
    (DM Sans Light reports itself as DM Sans, style Light), so it is never read from the files."""
    files: dict[tuple[str, bool, bool], Path] = {}
    for style in TYPOGRAPHY["styles"].values():
        if style["file"]:
            key = (style["family"], style["bold"], style.get("italic", False))
            files.setdefault(key, FONTS / style["file"])
    return files


@cache
def _font(family: str, bold: bool, italic: bool) -> tuple[ImageFont.FreeTypeFont, bool]:
    """The file a run is drawn with, loaded at 1000 px, and whether it is a family the
    typography names."""
    files = _font_files()
    named = family in {f for f, _, _ in files} or family in STAND_INS
    family = STAND_INS.get(family, family if named else DEFAULT_FAMILY)
    path = (
        files.get((family, bold, italic))
        or files.get((family, bold, False))
        or files.get((family, False, False))
        or files[(DEFAULT_FAMILY, False, False)]
    )
    return ImageFont.truetype(str(path), 1000), named


# --- laying text out --------------------------------------------------------------------


@dataclass
class Piece:
    text: str
    width: float  # inches
    size: float  # pt


@dataclass
class Paragraph:
    lines: list[list[Piece]]  # one per hard line (a:br)
    size: float  # pt, the height of an empty paragraph
    align: Any
    spacing: float  # multiple of single spacing
    spacing_pt: float | None  # exact pitch, when set in points
    before: float  # inches
    after: float  # inches
    margin: float  # inches from the frame's text edge to the text
    first: float  # the same, for the first line (the indent applied)


def _paragraphs(text_frame: Any, report: Report, where: str) -> list[Paragraph]:
    out = []
    for p in text_frame.paragraphs:
        ppr = p._p.find(qn("a:pPr"))
        margin = int(ppr.get("marL", 0)) / EMU if ppr is not None else 0.0
        indent = int(ppr.get("indent", 0)) / EMU if ppr is not None else 0.0
        # a bullet hangs in the indent; the text of every line starts at the margin
        bulleted = ppr is not None and any(
            ppr.find(qn(t)) is not None for t in ("a:buChar", "a:buAutoNum")
        )
        lines: list[list[Piece]] = [[]]
        sizes = []
        for child in p._p:
            tag = child.tag.split("}")[1]
            if tag == "br":
                lines.append([])
                continue
            if tag not in ("r", "fld"):
                continue
            rpr = child.find(qn("a:rPr"))
            t = child.find(qn("a:t"))
            s = t.text if t is not None and t.text else ""
            get = rpr.get if rpr is not None else (lambda _k: None)
            size = int(get("sz")) / 100 if get("sz") else DEFAULT_SIZE
            latin = rpr.find(qn("a:latin")) if rpr is not None else None
            family = latin.get("typeface") if latin is not None else DEFAULT_FAMILY
            font, named = _font(family, get("b") in ("1", "true"), get("i") in ("1", "true"))
            if not named and s.strip():
                report.note(f"{where}: {family!r} is not a shipped font; measured as DM Sans")
            sizes.append(size)
            for chunk in re.findall(r"\s+|\S+", s):
                lines[-1].append(Piece(chunk, font.getlength(chunk) / 1000 * size / 72, size))
        end = p._p.find(qn("a:endParaRPr"))
        spacing, spacing_pt = 1.0, None
        if isinstance(p.line_spacing, Length):
            spacing_pt = p.line_spacing.pt
        elif p.line_spacing is not None:
            spacing = float(p.line_spacing)
        out.append(
            Paragraph(
                lines=lines,
                size=max(sizes)
                if sizes
                else (
                    int(end.get("sz")) / 100 if end is not None and end.get("sz") else DEFAULT_SIZE
                ),
                align=p.alignment,
                spacing=spacing,
                spacing_pt=spacing_pt,
                before=p.space_before.inches if p.space_before is not None else 0.0,
                after=p.space_after.inches if p.space_after is not None else 0.0,
                margin=margin,
                first=margin if bulleted else margin + indent,
            )
        )
    return out


def _wrap(pieces: list[Piece], room_first: float, room: float) -> tuple[list[list[Piece]], bool]:
    """Wrap at spaces as PowerPoint does; a word wider than the line is broken between
    characters. Returns the lines and whether there was such a word."""
    words: list[list[Piece]] = []  # pieces joined up to the next space; spaces on their own
    for piece in pieces:
        if not piece.text.isspace() and words and not words[-1][-1].text.isspace():
            words[-1].append(piece)
        else:
            words.append([piece])
    lines: list[list[Piece]] = [[]]
    width = 0.0  # including trailing spaces, which never force a break
    too_wide = False

    def room_now() -> float:
        return room_first if len(lines) == 1 else room

    for word in words:
        w = sum(p.width for p in word)
        if word[0].text.isspace():
            lines[-1] += word
            width += w
            continue
        if _ink(lines[-1]) and width + w > room_now() + 1e-6:
            lines.append([])
            width = 0.0
        if w <= room_now() + 1e-6:
            lines[-1] += word
            width += w
            continue
        too_wide = True
        for piece in word:
            per_char = piece.width / max(1, len(piece.text))
            for ch in piece.text:
                if width > 0 and width + per_char > room_now() + 1e-6:
                    lines.append([])
                    width = 0.0
                lines[-1].append(Piece(ch, per_char, piece.size))
                width += per_char
    return lines, too_wide


def _ink(line: list[Piece]) -> list[Piece]:
    """The line without its trailing spaces."""
    end = len(line)
    while end and line[end - 1].text.isspace():
        end -= 1
    return line[:end]


@dataclass
class Layout:
    lines: list[Line]
    needed: float  # inches of height the text takes
    pitch: float  # the largest line pitch, inches
    too_wide: bool


def _layout(
    text_frame: Any,
    box: tuple[float, float, float, float],
    insets: tuple[float, float, float, float],
    anchor: Any,
    slide: int,
    name: str,
    order: int,
    report: Report,
) -> Layout:
    x, y, w, h = box
    left, right, top, bottom = insets
    width = w - left - right
    wrap = text_frame.word_wrap is not False
    # (text, width, left edge, room, size, pitch, space above, align) per line
    rows: list[tuple[str, float, float, float, float, float, float, Any]] = []
    too_wide = False
    paragraphs = _paragraphs(text_frame, report, f"slide {slide}, {name}")
    for k, para in enumerate(paragraphs):
        first = True
        for hard in para.lines:
            if wrap:
                wrapped, wide = _wrap(
                    hard, width - (para.first if first else para.margin), width - para.margin
                )
                too_wide |= wide
            else:
                wrapped = [hard]
            for line in wrapped:
                kept = _ink(line)
                size = max((p.size for p in kept), default=para.size)
                pitch = (
                    para.spacing_pt / 72
                    if para.spacing_pt is not None
                    else LINE_PITCH * size * para.spacing / 72
                )
                start = para.first if first else para.margin
                above = (para.before + (paragraphs[k - 1].after if k else 0.0)) if first else 0.0
                text = "".join(p.text for p in kept).strip()
                rows.append(
                    (
                        text,
                        sum(p.width for p in kept),
                        start,
                        width - start,
                        size,
                        pitch,
                        above,
                        para.align,
                    )
                )
                first = False
    needed = sum(r[5] + r[6] for r in rows)
    inner = h - top - bottom
    if text_frame.auto_size == MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT or anchor in (None, MSO_ANCHOR.TOP):
        cursor = y + top
    elif anchor == MSO_ANCHOR.MIDDLE:
        cursor = y + top + (inner - needed) / 2
    else:
        cursor = y + h - bottom - needed
    lines = []
    for text, lw, start, room, size, pitch, above, align in rows:
        cursor += above
        if text:
            shift = 0.0
            if align == PP_ALIGN.CENTER:
                shift = (room - lw) / 2
            elif align == PP_ALIGN.RIGHT:
                shift = room - lw
            x0 = x + left + start + shift
            lines.append(
                Line(slide, name, order, text, x0, x0 + lw, cursor + BASELINE * pitch, size)
            )
        cursor += pitch
    return Layout(lines, needed, max((r[5] for r in rows), default=0.0), too_wide)


# --- reading the slides -----------------------------------------------------------------

Transform = tuple[float, float, float, float]  # slide inches = emu * scale + offset, per axis


def _place(shape: Any, xf: Transform) -> tuple[float, float, float, float]:
    sx, ox, sy, oy = xf
    return (
        shape.left * sx + ox,
        shape.top * sy + oy,
        shape.width * sx,
        shape.height * sy,
    )


def _child_transform(group: Any, xf: Transform) -> Transform:
    sx, ox, sy, oy = xf
    xfrm = group._element.grpSpPr.find(qn("a:xfrm"))
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    ch_off, ch_ext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
    kx = int(ext.get("cx")) / max(1, int(ch_ext.get("cx")))
    ky = int(ext.get("cy")) / max(1, int(ch_ext.get("cy")))
    return (
        sx * kx,
        ox + sx * (int(off.get("x")) - int(ch_off.get("x")) * kx),
        sy * ky,
        oy + sy * (int(off.get("y")) - int(ch_off.get("y")) * ky),
    )


def _filled(shape: Any) -> bool:
    sppr = shape._element.find(qn("p:spPr"))
    return sppr is not None and any(
        sppr.find(qn(t)) is not None
        for t in ("a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill")
    )


def _insets(frame: Any) -> tuple[float, float, float, float]:
    return (
        frame.margin_left / EMU,
        frame.margin_right / EMU,
        frame.margin_top / EMU,
        frame.margin_bottom / EMU,
    )


def _snippet(lines: list[Line], fallback: str = "") -> str:
    text = lines[0].text if lines else fallback
    return text if len(text) <= 40 else text[:39] + "…"


def _read_slide(slide: Any, number: int, size: tuple[float, float], report: Report) -> Slide:
    out = Slide(number)
    order = 0

    def overflow(
        name: str, lay: Layout, box: tuple[float, float, float, float], inner: float, fit: bool
    ) -> None:
        x, y, w, h = box
        if not fit and lay.needed - inner > OVERFLOW_SHARE * lay.pitch:
            report.findings.append(
                Finding(
                    number,
                    "text overflowing its box",
                    f"{name} {_snippet(lay.lines)!r}: needs {lay.needed:.2f} in, "
                    f"the box holds {inner:.2f} in",
                    _r(x, y, x + w, y + h),
                )
            )
        if lay.too_wide:
            report.findings.append(
                Finding(
                    number,
                    "text overflowing its box",
                    f"{name} {_snippet(lay.lines)!r}: a word is wider than the box and is "
                    "broken mid-word",
                    _r(x, y, x + w, y + h),
                )
            )

    def walk(shapes: Any, xf: Transform) -> None:
        nonlocal order
        for shape in shapes:
            order += 1
            tag = shape._element.tag.split("}")[1]
            name = shape.name
            if tag == "grpSp":
                walk(shape.shapes, _child_transform(shape, xf))
                continue
            x, y, w, h = _place(shape, xf)
            if getattr(shape, "rotation", 0):
                report.note(f"slide {number}, {name}: rotated, not checked")
                continue
            if tag == "cxnSp":
                sx, ox, sy, oy = xf
                out.segments.append(
                    Segment(
                        name,
                        shape.begin_x * sx + ox,
                        shape.begin_y * sy + oy,
                        shape.end_x * sx + ox,
                        shape.end_y * sy + oy,
                    )
                )
                continue
            if tag == "pic":
                background = order == 1 and w * h >= BACKGROUND_SHARE * size[0] * size[1]
                if not background:
                    out.pictures.append(Rect(name, order, x, y, x + w, y + h))
                continue
            if tag == "graphicFrame":
                if getattr(shape, "has_table", False) and shape.has_table:
                    _read_table(shape, (x, y, w, h), number, order, out, report)
                else:  # a chart or another object: text over it collides with it
                    out.pictures.append(Rect(name, order, x, y, x + w, y + h))
                continue
            if tag != "sp":
                continue
            if _filled(shape):
                out.fills.append(Rect(name, order, x, y, x + w, y + h))
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            tf = shape.text_frame
            insets = _insets(tf)
            lay = _layout(tf, (x, y, w, h), insets, tf.vertical_anchor, number, name, order, report)
            out.lines += lay.lines
            fit = tf.auto_size == MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
            overflow(name, lay, (x, y, w, h), h - insets[2] - insets[3], fit)

    def _r(*v: float) -> tuple[float, float, float, float]:
        return tuple(round(c, 2) for c in v)  # type: ignore[return-value]

    xf = (1 / EMU, 0.0, 1 / EMU, 0.0)
    # What the slide's layout draws under it (background, cards, rules, fixed text),
    # then the slide's own shapes on top. A layout's placeholders are not drawn: the slide
    # has its own copy of each one it uses.
    walk([sh for sh in slide.slide_layout.shapes if not sh.is_placeholder], xf)
    walk(slide.shapes, xf)
    return out


def _read_table(
    frame: Any,
    box: tuple[float, float, float, float],
    number: int,
    order: int,
    out: Slide,
    report: Report,
) -> None:
    """Rows grow to fit their cells, as PowerPoint draws them; the rows below move down."""
    table = frame.table
    x, y, _, _ = box
    col_w = [c.width / EMU for c in table.columns]
    col_x = [x + sum(col_w[:i]) for i in range(len(col_w))]
    top = y
    for r, row in enumerate(table.rows):
        set_h = row.height / EMU
        laid = []
        grown = set_h
        for c, cell in enumerate(row.cells):
            if cell.is_spanned:
                continue
            cw = sum(col_w[c : c + cell.span_width])
            insets = _insets(cell)
            anchor = cell.vertical_anchor
            name = f"{frame.name} cell ({r + 1}, {c + 1})"
            probe = _layout(
                cell.text_frame,
                (0, 0, cw, set_h),
                insets,
                MSO_ANCHOR.TOP,
                number,
                name,
                order,
                report,
            )
            if cell.span_height == 1:
                grown = max(grown, probe.needed + insets[2] + insets[3])
            if probe.too_wide:
                report.findings.append(
                    Finding(
                        number,
                        "text overflowing its box",
                        f"{name} {_snippet(probe.lines)!r}: a word is wider than the column",
                        (
                            round(col_x[c], 2),
                            round(top, 2),
                            round(col_x[c] + cw, 2),
                            round(top + set_h, 2),
                        ),
                    )
                )
            laid.append((cell, c, cw, insets, anchor, name))
        if grown - set_h > TOLERANCE:
            report.findings.append(
                Finding(
                    number,
                    "table row grown",
                    f"{frame.name} row {r + 1}: its text needs {grown:.2f} in, the row is set to "
                    f"{set_h:.2f}; PowerPoint makes it taller and moves every row below down by "
                    f"{grown - set_h:.2f} in",
                    (round(x, 2), round(top, 2), round(x + sum(col_w), 2), round(top + grown, 2)),
                )
            )
        for cell, c, cw, insets, anchor, name in laid:
            h = (
                grown
                if cell.span_height == 1
                else sum(
                    table.rows[i].height / EMU
                    for i in range(r, min(len(table.rows), r + cell.span_height))
                )
            )
            lay = _layout(
                cell.text_frame, (col_x[c], top, cw, h), insets, anchor, number, name, order, report
            )
            out.lines += lay.lines
        top += grown


# --- the checks -------------------------------------------------------------------------


def _overlap(a0: float, a1: float, b0: float, b1: float) -> float:
    return min(a1, b1) - max(a0, b0)


def _hits(line: Line, rect: Rect) -> bool:
    return (
        _overlap(line.x0, line.x1, rect.x0, rect.x1) > TOLERANCE
        and _overlap(line.top, line.bottom, rect.y0, rect.y1) > TOLERANCE
    )


def _inside(line: Line, rect: Rect) -> bool:
    return (
        line.x0 >= rect.x0 - TOLERANCE
        and line.x1 <= rect.x1 + TOLERANCE
        and line.top >= rect.y0 - TOLERANCE
        and line.bottom <= rect.y1 + TOLERANCE
    )


def _box(line: Line) -> tuple[float, float, float, float]:
    return (round(line.x0, 2), round(line.top, 2), round(line.x1, 2), round(line.bottom, 2))


def _short(text: str, n: int = 40) -> str:
    return text if len(text) <= n else text[: n - 1] + "…"


def _check(s: Slide, size: tuple[float, float]) -> list[Finding]:
    found: dict[tuple[str, ...], Finding] = {}  # one finding per pair of things

    def add(key: tuple[str, ...], kind: str, detail: str, line: Line) -> None:
        found.setdefault((kind, *key), Finding(s.number, kind, detail, _box(line)))

    ordered = sorted(s.lines, key=lambda ln: ln.top)
    for i, a in enumerate(ordered):
        for b in ordered[i + 1 :]:
            if b.top >= a.bottom - TOLERANCE:
                break
            if _overlap(a.x0, a.x1, b.x0, b.x1) > TOLERANCE:
                pair = tuple(sorted((a.shape, b.shape)))
                add(
                    pair,
                    "text on text",
                    f"{a.shape} {_short(a.text)!r} and {b.shape} {_short(b.text)!r}",
                    a,
                )
    for line in s.lines:
        for pic in s.pictures:
            if _hits(line, pic):
                add(
                    (line.shape, pic.name),
                    "text on a picture",
                    f"{line.shape} {_short(line.text)!r} on {pic.name}",
                    line,
                )
        for seg in s.segments:
            horizontal = abs(seg.y0 - seg.y1) < 0.005
            vertical = abs(seg.x0 - seg.x1) < 0.005
            if horizontal:
                # through the letters, from the cap height to just above the baseline, so an
                # underline below the baseline is not one
                through = line.top + TOLERANCE < seg.y0 < line.baseline - 0.1 * line.size / 72
                crosses = (
                    through
                    and _overlap(line.x0, line.x1, min(seg.x0, seg.x1), max(seg.x0, seg.x1))
                    > TOLERANCE
                )
            elif vertical:
                inside = line.x0 + TOLERANCE < seg.x0 < line.x1 - TOLERANCE
                crosses = (
                    inside
                    and _overlap(line.top, line.bottom, min(seg.y0, seg.y1), max(seg.y0, seg.y1))
                    > TOLERANCE
                )
            else:
                crosses = False
            if crosses:
                add(
                    (line.shape, seg.name),
                    "text crossed by a line",
                    f"{line.shape} {_short(line.text)!r} by {seg.name}",
                    line,
                )
        for fill in s.fills:
            if fill.name == line.shape or not _hits(line, fill):
                continue
            if fill.order > line.order:
                add(
                    (line.shape, fill.name),
                    "text under a shape",
                    f"{line.shape} {_short(line.text)!r} is covered by {fill.name}",
                    line,
                )
            elif not _inside(line, fill):
                add(
                    (line.shape, fill.name),
                    "text out of its card",
                    f"{line.shape} {_short(line.text)!r} runs off {fill.name}",
                    line,
                )
        width, height = size
        if (
            line.x0 < -TOLERANCE
            or line.x1 > width + TOLERANCE
            or line.top < -TOLERANCE
            or line.bottom > height + TOLERANCE
        ):
            add((line.shape,), "text off the slide", f"{line.shape} {_short(line.text)!r}", line)
        elif abs(width - 10) < 0.01 and (
            line.x0 < SAFE_X[0] - SAFE_TOLERANCE or line.x1 > SAFE_X[1] + SAFE_TOLERANCE
        ):
            add(
                (line.shape,),
                "text outside the safe area",
                f"{line.shape} {_short(line.text)!r}",
                line,
            )
    return list(found.values())


def check(path: Path | str) -> Report:
    """Every collision in the deck."""
    prs = Presentation(str(path))
    size = (prs.slide_width / EMU, prs.slide_height / EMU)
    report = Report()
    for number, slide in enumerate(prs.slides, start=1):
        read = _read_slide(slide, number, size, report)
        report.findings += _check(read, size)
    report.findings.sort(key=lambda f: (f.slide, f.where[1], f.where[0]))
    return report


@dataclass
class Fit:
    """How the text of one box fits it, in inches."""

    needed: float
    holds: float
    pitch: float  # the largest line pitch
    chars: int  # characters in the box now
    too_wide: bool  # a word wider than the box

    @property
    def ok(self) -> bool:
        return not self.too_wide and self.needed - self.holds <= OVERFLOW_SHARE * self.pitch

    @property
    def chars_that_fit(self) -> int:
        """About how many characters the box holds, at the text's size and spacing."""
        if self.needed <= 0:
            return self.chars
        return max(0, int(self.chars * self.holds / self.needed))


def fit(shape: Any) -> Fit:
    """How the text of one text box fits it, measured as `check` measures it, for build
    code that refuses a box the moment it is written rather than after the deck is saved."""
    tf = shape.text_frame
    box = (shape.left / EMU, shape.top / EMU, shape.width / EMU, shape.height / EMU)
    insets = _insets(tf)
    lay = _layout(tf, box, insets, tf.vertical_anchor, 0, shape.name, 0, Report())
    holds = box[3] - insets[2] - insets[3]
    if tf.auto_size == MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT:
        holds = max(holds, lay.needed)
    chars = sum(len(p.text) for p in tf.paragraphs)
    return Fit(lay.needed, holds, lay.pitch, chars, lay.too_wide)


# --- rendering --------------------------------------------------------------------------


def _soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        if found := shutil.which(name):
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def render(path: Path, out_dir: Path) -> list[Path]:
    """The deck as a PDF, and as one PNG per slide if pymupdf is installed, in out_dir."""
    soffice = _soffice()
    if soffice is None:
        raise RuntimeError("LibreOffice (soffice) is not installed; nothing rendered")
    out_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as profile:
        fonts = Path(profile) / "user" / "fonts"
        fonts.mkdir(parents=True)
        for ttf in FONTS.glob("*.ttf"):
            shutil.copy(ttf, fonts)
        subprocess.run(
            [
                soffice,
                f"-env:UserInstallation={Path(profile).as_uri()}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(out_dir),
                str(path),
            ],
            check=True,
            capture_output=True,
            timeout=300,
        )
    pdf = out_dir / f"{path.stem}.pdf"
    made = [pdf]
    try:
        import pymupdf
    except ImportError:
        return made
    with pymupdf.open(pdf) as doc:
        for i, page in enumerate(doc, start=1):
            png = out_dir / f"{path.stem}-{i:02d}.png"
            page.get_pixmap(dpi=110).save(png)
            made.append(png)
    return made


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--json", action="store_true", help="print the findings as JSON")
    parser.add_argument(
        "--render", type=Path, metavar="DIR", help="also render the deck to a PDF and PNGs in DIR"
    )
    args = parser.parse_args(argv)
    report = check(args.pptx)
    if args.json:
        print(
            json.dumps(
                {"findings": [asdict(f) for f in report.findings], "notes": report.notes},
                indent=2,
                ensure_ascii=False,
            )
        )
    else:
        if not report.findings:
            print(f"{args.pptx}: no collisions")
        else:
            print(f"{args.pptx}: {len(report.findings)} collision(s)")
            for f in report.findings:
                x0, y0, x1, y1 = f.where
                print(f"  slide {f.slide}: {f.kind}: {f.detail}  at ({x0}, {y0})–({x1}, {y1}) in")
        for note in report.notes:
            print(f"  note: {note}")
    if args.render:
        try:
            for made in render(args.pptx, args.render):
                print(f"rendered: {made}")
        except (RuntimeError, subprocess.SubprocessError) as e:
            print(f"render failed: {e}", file=sys.stderr)
    return 1 if report.findings else 0


if __name__ == "__main__":
    sys.exit(main())
